"""Recompute laser features for stored scans from their frames on disk.

Features are a derived product: whenever the computation changes (gain
normalization, a new metric) every scan can be brought up to date without
touching the bench. Channels that were flagged invalid (``valid: false``,
e.g. a laser that did not fire) stay invalid, with their new numbers nulled
the same way.

    python -m scanner_system.recompute_features            # all scans
    python -m scanner_system.recompute_features --since 2026-09-15
    python -m scanner_system.recompute_features --sheet collection.xlsx

With --sheet, the workbook's laser_features tab is rebuilt from the
database so it matches (the tab is otherwise append-only).
"""

from __future__ import annotations

import argparse
import os
import sys
from typing import Any

from . import scanner_db, schema
from .capture import STORAGE_ROOT
from .laser_features import compute_features, feature_rows


def _invalidate(new_ch: dict[str, Any], old_ch: dict[str, Any]) -> dict[str, Any]:
    nulled = {k: (None if isinstance(v, (int, float)) and not isinstance(v, bool)
                  else (False if isinstance(v, bool) else v)) for k, v in new_ch.items()}
    nulled.update({"valid": False, "invalid_reason": old_ch.get("invalid_reason"),
                   "raw_invalid": {k: v for k, v in new_ch.items()}})
    return nulled


def recompute(since: str = "", db=None, storage_root: str = STORAGE_ROOT) -> dict[str, int]:
    d = scanner_db.get_db(db)
    q: dict[str, Any] = {"status": {"$in": ["complete", "partial", "superseded"]}}
    if since:
        q["started_at"] = {"$gte": since}
    done, skipped = 0, 0
    for scan in d["scans"].find(q):
        laser_dir = os.path.join(storage_root, "scans", scan["_id"], "laser")
        feats = compute_features(laser_dir)
        if not feats:
            skipped += 1
            continue
        old = (scan.get("laser_features") or {}).get("channels") or {}
        for ch, oc in old.items():
            if oc.get("valid") is False and ch in feats["channels"]:
                feats["channels"][ch] = _invalidate(feats["channels"][ch], oc)
        if any(c.get("valid") is False for c in feats["channels"].values()):
            feats["red_green_ratio"] = None
            feats["red2_green_ratio"] = None
        scanner_db.set_laser_features(scan["_id"], feats, db=d)

        # Fringe features against the reference the scan was captured with
        # (recorded in capture_config; a reference from another geometry
        # would be wrong, so scans without the record keep what they have).
        ref_name = (scan.get("capture_config") or {}).get("reference_dir")
        fringe_dir = os.path.join(storage_root, "scans", scan["_id"], "fringe")
        if ref_name and os.path.isfile(os.path.join(fringe_dir, "scan.npz")):
            from .fringe_features import compute_fringe_features
            repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            ref_dir = os.path.join(repo, "data", "scan_test", "calib_new", ref_name)
            try:
                ff = compute_fringe_features(fringe_dir, ref_dir)
            except Exception:
                ff = None
            if ff:
                scanner_db.set_scan_meta(scan["_id"], {"fringe_features": ff}, db=d)
        done += 1
    return {"recomputed": done, "skipped_no_frames": skipped}


def rebuild_sheet_features(path: str, db=None) -> int:
    from openpyxl import load_workbook
    from . import manifest

    wb = load_workbook(path)
    if manifest.FEATURES_SHEET in wb.sheetnames:
        del wb[manifest.FEATURES_SHEET]
        wb.save(path)
    rows: list[dict[str, Any]] = []
    for sample in scanner_db.list_samples(db=db):
        for scan in scanner_db.scans_for_sample(sample["_id"], db=db):
            if scan.get("status") == "complete" and scan.get("laser_features"):
                rows.extend(feature_rows(sample, scan, schema.LASER_WAVELENGTHS_NM))
    rows.sort(key=lambda r: (r.get("label") or "", r.get("angle_index") or 0, r.get("channel") or 0))
    manifest.write_features(path, rows)
    return len(rows)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--since", default="", help="only scans started on/after this ISO date")
    ap.add_argument("--sheet", default=None, help="also rebuild this workbook's laser_features tab")
    args = ap.parse_args(argv)
    out = recompute(since=args.since)
    print(f"recomputed {out['recomputed']} scans ({out['skipped_no_frames']} had no usable frames)")
    if args.sheet:
        n = rebuild_sheet_features(args.sheet)
        print(f"rebuilt laser_features tab in {args.sheet}: {n} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
